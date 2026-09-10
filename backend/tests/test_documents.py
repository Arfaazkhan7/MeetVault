import io
import pytest
import pypdf
import pptx
from pptx.util import Inches

def register_user(client, email="docuser@example.com", name="Doc User", password="Password123!"):
    r = client.post("/api/v1/auth/register", json={
        "full_name": name, "email": email, "password": password,
    })
    assert r.status_code == 201, r.text
    return r.json()

def login_user(client, email="docuser@example.com", password="Password123!"):
    r = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]

def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}

def create_sample_pdf_bytes() -> bytes:
    """Create in-memory PDF with 2 pages."""
    writer = pypdf.PdfWriter()
    page1 = writer.add_blank_page(width=612, height=792)
    page2 = writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

def create_sample_pptx_bytes() -> bytes:
    """Create in-memory PPTX with 2 slides containing title and body text."""
    prs = pptx.Presentation()
    
    # Slide 1
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]
    title1.text = "Quarterly Business Review"
    subtitle1.text = "Q3 Strategy Presentation"

    # Slide 2
    slide_layout2 = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout2)
    title2 = slide2.shapes.title
    body2 = slide2.placeholders[1]
    title2.text = "Revenue Projections"
    body2.text = "Targeting 25 percent growth in Enterprise sales."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def setup_doc_context(client):
    user = register_user(client)
    token = login_user(client)
    comp = client.post("/api/v1/companies", json={"name": "DocCorp"}, headers=auth_headers(token)).json()
    cid = comp["public_id"]

    proj = client.post(f"/api/v1/companies/{cid}/projects", json={"name": "DocProject"}, headers=auth_headers(token)).json()
    pid = proj["public_id"]

    meet = client.post(f"/api/v1/companies/{cid}/meetings", json={"title": "DocMeeting", "project_id": pid}, headers=auth_headers(token)).json()
    mid = meet["public_id"]

    return token, cid, pid, mid


def test_upload_validation(client):
    token, cid, _, _ = setup_doc_context(client)
    
    # Unsupported file type (.txt)
    files = {"file": ("notes.txt", b"Unsupported content", "text/plain")}
    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files=files, headers=auth_headers(token))
    assert r.status_code == 400
    assert "Only PDF, PPT, and PPTX" in r.json()["detail"]

    # Unsupported file type (.exe)
    files = {"file": ("script.exe", b"binary content", "application/octet-stream")}
    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files=files, headers=auth_headers(token))
    assert r.status_code == 400


def test_pdf_processing(client):
    token, cid, pid, mid = setup_doc_context(client)
    pdf_bytes = create_sample_pdf_bytes()

    files = {"file": ("report.pdf", pdf_bytes, "application/pdf")}
    data = {"title": "Annual Financial Report", "project_id": pid, "meeting_id": mid}

    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files=files, data=data, headers=auth_headers(token))
    assert r.status_code == 201, r.text
    doc = r.json()

    assert doc["title"] == "Annual Financial Report"
    assert doc["file_type"] == "pdf"
    assert doc["page_count"] == 2
    assert doc["project_id"] == pid
    assert doc["meeting_id"] == mid
    assert doc["status"] == "PROCESSED"
    assert len(doc["pages"]) == 2
    assert doc["pages"][0]["page_number"] == 1
    assert doc["pages"][1]["page_number"] == 2


def test_pptx_processing(client):
    token, cid, pid, _ = setup_doc_context(client)
    pptx_bytes = create_sample_pptx_bytes()

    files = {"file": ("strategy.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    data = {"title": "Q3 Strategy Deck", "project_id": pid}

    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files=files, data=data, headers=auth_headers(token))
    assert r.status_code == 201, r.text
    doc = r.json()

    assert doc["title"] == "Q3 Strategy Deck"
    assert doc["file_type"] == "pptx"
    assert doc["page_count"] == 2
    assert doc["project_id"] == pid

    # Verify extracted text on slides
    pages = doc["pages"]
    assert len(pages) == 2
    assert "Quarterly Business Review" in pages[0]["text_content"]
    assert "Revenue Projections" in pages[1]["text_content"] or "25 percent" in pages[1]["text_content"]


def test_ppt_processing(client):
    token, cid, _, _ = setup_doc_context(client)
    ppt_bytes = b"PowerPoint Document binary stream mock data with Slide 1 title and Slide 2 notes"

    files = {"file": ("legacy.ppt", ppt_bytes, "application/vnd.ms-powerpoint")}
    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files=files, headers=auth_headers(token))
    assert r.status_code == 201, r.text
    doc = r.json()

    assert doc["file_type"] == "ppt"
    assert doc["page_count"] >= 1
    assert doc["status"] == "PROCESSED"


def test_document_listing_and_search(client):
    token, cid, pid, _ = setup_doc_context(client)

    # Upload PDF
    pdf_bytes = create_sample_pdf_bytes()
    client.post(f"/api/v1/companies/{cid}/documents/upload", files={"file": ("alpha.pdf", pdf_bytes, "application/pdf")}, data={"title": "Alpha Spec"}, headers=auth_headers(token))

    # Upload PPTX
    pptx_bytes = create_sample_pptx_bytes()
    client.post(f"/api/v1/companies/{cid}/documents/upload", files={"file": ("beta.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, data={"title": "Beta Overview", "project_id": pid}, headers=auth_headers(token))

    # List all documents
    r = client.get(f"/api/v1/companies/{cid}/documents", headers=auth_headers(token))
    assert r.status_code == 200
    docs = r.json()
    assert len(docs) == 2

    # Filter by project
    r_proj = client.get(f"/api/v1/companies/{cid}/documents?project_id={pid}", headers=auth_headers(token))
    assert r_proj.status_code == 200
    assert len(r_proj.json()) == 1
    assert r_proj.json()[0]["title"] == "Beta Overview"

    # Full text search in slide text ("Revenue")
    r_search = client.get(f"/api/v1/companies/{cid}/documents/search?q=Revenue", headers=auth_headers(token))
    assert r_search.status_code == 200
    search_res = r_search.json()
    assert search_res["total"] >= 1
    matched_doc = search_res["results"][0]["document"]
    assert matched_doc["title"] == "Beta Overview"


def test_document_association(client):
    token, cid, pid, mid = setup_doc_context(client)
    pdf_bytes = create_sample_pdf_bytes()

    # Upload unassociated document
    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files={"file": ("standalone.pdf", pdf_bytes, "application/pdf")}, headers=auth_headers(token))
    doc = r.json()
    doc_id = doc["public_id"]
    assert doc["project_id"] is None
    assert doc["meeting_id"] is None

    # Associate with Meeting and Project
    r_assoc = client.patch(
        f"/api/v1/companies/{cid}/documents/{doc_id}/associate",
        json={"project_id": pid, "meeting_id": mid},
        headers=auth_headers(token)
    )
    assert r_assoc.status_code == 200
    updated_doc = r_assoc.json()
    assert updated_doc["project_id"] == pid
    assert updated_doc["meeting_id"] == mid

    # Get document details to verify Meeting ↔ Document ↔ Page/Slide hierarchy
    r_detail = client.get(f"/api/v1/companies/{cid}/documents/{doc_id}", headers=auth_headers(token))
    assert r_detail.status_code == 200
    detail = r_detail.json()
    assert detail["meeting_id"] == mid
    assert len(detail["pages"]) == 2


def test_private_storage_and_download(client):
    token, cid, _, _ = setup_doc_context(client)
    pdf_bytes = create_sample_pdf_bytes()

    r = client.post(f"/api/v1/companies/{cid}/documents/upload", files={"file": ("private_doc.pdf", pdf_bytes, "application/pdf")}, headers=auth_headers(token))
    doc = r.json()
    doc_id = doc["public_id"]

    # Download file via authenticated API
    r_dl = client.get(f"/api/v1/companies/{cid}/documents/{doc_id}/download", headers=auth_headers(token))
    assert r_dl.status_code == 200
    assert r_dl.content == pdf_bytes
