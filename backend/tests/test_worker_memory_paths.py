"""Exercise the lightweight Worker paths against real tokens and OOXML files."""
import asyncio
import base64
import io
import json
import shutil
import subprocess
import sys
import time
import zipfile
from types import SimpleNamespace

import httpx
import pptx
import pytest
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import rsa
from cryptography.x509.oid import NameOID
from datetime import datetime, timedelta, timezone
from jose import JWTError, jwt

from app.core.config import settings
from app.core.workers_crypto import certificate_spki
from app.services import document_service, firebase_auth_service
from app.services.pptx_text import extract_slides


@pytest.fixture
def webcrypto(monkeypatch):
    """Bridge the Python adapter to Node's actual Web Crypto implementation.

    This validates key format, algorithm, byte order and signature behavior;
    actual JSPI/Pyodide integration still requires a Worker runtime test.
    """
    if not shutil.which("node"):
        pytest.skip("Node required to test Web Crypto")

    class Subtle:
        async def importKey(self, kind, data, algorithm, extractable, usages):
            assert (kind, algorithm, extractable, usages) == (
                "spki", {"name": "RSASSA-PKCS1-v1_5", "hash": "SHA-256"}, False, ["verify"],
            )
            return data

        async def verify(self, algorithm, key, signature, message):
            assert algorithm == "RSASSA-PKCS1-v1_5"
            encoded = [base64.b64encode(v).decode() for v in (key, signature, message)]
            script = """const {webcrypto} = require('node:crypto');
const fs = require('node:fs');
(async () => {
  const [key, signature, message] = JSON.parse(fs.readFileSync(0, 'utf8'))
    .map(value => Buffer.from(value, 'base64'));
  const imported = await webcrypto.subtle.importKey('spki', key,
    {name: 'RSASSA-PKCS1-v1_5', hash: 'SHA-256'}, false, ['verify']);
  console.log(JSON.stringify(await webcrypto.subtle.verify(
    'RSASSA-PKCS1-v1_5', imported, signature, message)));
})().catch(err => {console.error(err); process.exit(1);});
"""
            result = subprocess.run(["node", "-e", script], input=json.dumps(encoded),
                                    text=True, capture_output=True, check=True)
            return json.loads(result.stdout)

    monkeypatch.setitem(sys.modules, "js", SimpleNamespace(
        crypto=SimpleNamespace(subtle=Subtle()), Object=SimpleNamespace(fromEntries=object()),
    ))
    monkeypatch.setitem(sys.modules, "pyodide.ffi", SimpleNamespace(
        run_sync=asyncio.run, to_js=lambda value, **kwargs: value,
    ))
    monkeypatch.setattr(firebase_auth_service, "IS_WORKERS", True)


@pytest.fixture
def signed_token(monkeypatch):
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    now = datetime.now(timezone.utc)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "synthetic-test")])
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name)
            .public_key(key.public_key()).serial_number(1)
            .not_valid_before(now - timedelta(days=1))
            .not_valid_after(now + timedelta(days=1)).sign(key, hashes.SHA256()))
    cert_pem = cert.public_bytes(serialization.Encoding.PEM).decode()
    monkeypatch.setattr(httpx, "get", lambda *args, **kwargs: httpx.Response(
        200, json={"synthetic-key": cert_pem}, request=httpx.Request("GET", "https://example.invalid"),
    ))
    def sign(**changes):
        timestamp = int(time.time())
        claims = {"sub": "firebase-user", "aud": settings.FIREBASE_PROJECT_ID,
                  "iss": f"https://securetoken.google.com/{settings.FIREBASE_PROJECT_ID}",
                  "iat": timestamp - 1, "auth_time": timestamp - 5, "exp": timestamp + 60}
        claims.update(changes)
        return jwt.encode(claims, key, algorithm="RS256", headers={"kid": "synthetic-key"})
    return sign, cert_pem, key


def test_certificate_public_key_extraction(signed_token):
    _, cert, key = signed_token
    expected = key.public_key().public_bytes(serialization.Encoding.DER,
                                             serialization.PublicFormat.SubjectPublicKeyInfo)
    assert certificate_spki(cert) == expected


def test_workers_firebase_uses_webcrypto_and_preserves_claim_checks(webcrypto, signed_token):
    sign, _, _ = signed_token
    assert firebase_auth_service._verify_portable_firebase_token(sign())["uid"] == "firebase-user"
    for changes in ({"aud": "other"}, {"iss": "https://attacker.invalid"}, {"exp": 1},
                    {"iat": 9999999999}, {"auth_time": 9999999999}, {"sub": ""}):
        with pytest.raises(JWTError):
            firebase_auth_service._verify_portable_firebase_token(sign(**changes))


def test_workers_firebase_rejects_tampered_signature(webcrypto, signed_token):
    sign, _, _ = signed_token
    header, payload, signature = sign().split(".")
    claims = jwt.get_unverified_claims(sign())
    claims["sub"] = "attacker"
    changed = base64.urlsafe_b64encode(json.dumps(claims).encode()).rstrip(b"=").decode()
    with pytest.raises(JWTError):
        firebase_auth_service._verify_portable_firebase_token(f"{header}.{changed}.{signature}")


def test_workers_firebase_rejects_algorithm_substitution(webcrypto, signed_token):
    token = jwt.encode({"sub": "attacker"}, "synthetic-secret", algorithm="HS256",
                       headers={"kid": "synthetic-key"})
    with pytest.raises(JWTError):
        firebase_auth_service._verify_portable_firebase_token(token)


def presentation_bytes():
    prs = pptx.Presentation()
    for title in ("First slide", "Second slide"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        paragraph = slide.placeholders[1].text_frame.paragraphs[0]
        paragraph.text = "Hello "
        paragraph.add_run().text = "world"
        slide.notes_slide.notes_text_frame.text = f"Speaker notes for {title}"
    # File names no longer match presentation order after rearrangement.
    ids = prs.slides._sldIdLst
    ids.insert(0, ids[1])
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def test_pptx_keeps_presentation_order_runs_and_notes(monkeypatch):
    monkeypatch.setattr(document_service, "IS_WORKERS", True)
    slides = document_service.extract_pptx_slides(presentation_bytes())
    assert len(slides) == 2
    assert slides[0]["page_number"] == 1
    assert slides[0]["text_content"].startswith("Second slide\nHello world")
    assert "Notes: Speaker notes for Second slide" in slides[0]["text_content"]
    assert slides[1]["text_content"].startswith("First slide")


def test_pptx_rejects_xml_expansion_and_oversized_parts():
    for data in (b'<!DOCTYPE p [<!ENTITY x "expanded">]><p>&x;</p>', b"x" * (2 * 1024 * 1024 + 1)):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("ppt/presentation.xml", data)
        with pytest.raises(ValueError):
            extract_slides(output.getvalue())


def test_worker_pptx_invalid_archive_is_client_error(monkeypatch):
    from fastapi import HTTPException
    monkeypatch.setattr(document_service, "IS_WORKERS", True)
    with pytest.raises(HTTPException) as error:
        document_service.extract_pptx_slides(b"invalid")
    assert error.value.status_code == 400
